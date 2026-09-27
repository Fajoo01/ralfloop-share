from __future__ import annotations

"""Deterministic domain helpers for the F-Gas installation form.

No network access lives here. Drive and WhatsApp are adapters in the MCP server;
this module only extracts, validates and renders bounded structured data.
"""

from dataclasses import dataclass
from datetime import datetime
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any, Mapping


MODEL_DEFAULTS: dict[str, dict[str, Any]] = {
    # Verified in the historical F-Gas workflow from the manufacturer's data.
    "2AMW42U4RGC": {"brand": "Hisense", "refrigerant": "R32", "charge_kg": "0,95"},
    # Bosch Climate 5000 M official product data: code 7733701932, R32, 1.1 kg,
    # fluorinated refrigerant circuit not hermetically sealed.
    "CL5000M 41/2 E": {
        "brand": "Bosch",
        "refrigerant": "R32",
        "charge_kg": "1,1",
        "compressor_count": "1",
        "hermetically_sealed": False,
        "equipment_type": "POMPA DI CALORE FISSA",
    },
}

MODEL_PRODUCT_CODES = {"CL5000M 41/2 E": "7733701932"}

REQUIRED_FOR_RENDER = (
    "client_name", "client_cf", "client_type",
    "brand", "model", "serial", "refrigerant", "charge_kg", "compressor_count",
    "fixed_leak_detection", "hermetically_sealed", "equipment_type",
    "install_address", "civic", "city", "province", "use_destination",
    "intervention_date", "sold_by_installer", "purchase_reference", "purchase_date",
    "installer_name", "installer_cf", "gas_recovered", "gas_added",
)

CF_RE = re.compile(r"\b[A-Z]{6}[0-9]{2}[A-Z][0-9]{2}[A-Z][0-9]{3}[A-Z]\b", re.I)
EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
PHONE_RE = re.compile(r"(?<!\d)(?:\+39\s*)?(?:3\d{2}[ .-]?\d{3}[ .-]?\d{4}|0\d{1,3}[ .-]?\d{5,8})(?!\d)")
DATE_RE = re.compile(r"\b([0-3]?\d[/-][01]?\d[/-](?:19|20)?\d{2})\b")
REFRIGERANT_RE = re.compile(r"\b(R(?:32|410A|407C|134A|290|600A|404A|454B|452B))\b", re.I)
KG_RE = re.compile(r"(?<!\d)(\d{1,3}(?:[.,]\d{1,3})?)\s*kg\b", re.I)


@dataclass(frozen=True)
class ValidationResult:
    ok: bool
    missing: tuple[str, ...]
    warnings: tuple[str, ...]
    normalized: dict[str, Any]


def normalize_record(record: Mapping[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for key, value in record.items():
        if isinstance(value, str):
            value = " ".join(value.split()).strip()
        out[key] = value
    for key in ("client_cf", "installer_cf", "serial", "refrigerant", "province"):
        if isinstance(out.get(key), str):
            out[key] = out[key].upper()
    if isinstance(out.get("charge_kg"), str):
        charge = out["charge_kg"].replace(".", ",")
        if "," in charge:
            whole, fraction = charge.split(",", 1)
            fraction = fraction.rstrip("0")
            charge = whole if not fraction else f"{whole},{fraction}"
        out["charge_kg"] = charge
    model = str(out.get("model") or "").upper().strip()
    if model:
        out["model"] = model
        defaults = MODEL_DEFAULTS.get(model, {})
        for key, value in defaults.items():
            out.setdefault(key, value)
            if not out.get(key):
                out[key] = value
    return out


def extract_record(text: str, overrides: Mapping[str, Any] | None = None) -> dict[str, Any]:
    text = text or ""
    record: dict[str, Any] = {}

    def label(*names: str) -> str:
        for name in names:
            m = re.search(rf"(?im)^\s*{name}\s*[:\-]\s*([^\n\r]{{1,180}})", text)
            if m:
                return m.group(1).strip()
        return ""

    record["client_name"] = label(r"(?:cliente|nome\s+cliente|nominativo)")
    record["client_email"] = label(r"(?:mail|e-?mail)(?:\s+cliente)?")
    record["client_phone"] = label(r"(?:telefono\s+cliente|tel\.?\s+cliente|cellulare\s+cliente|cell\.?\s+cliente)")
    record["install_address"] = label(r"(?:indirizzo|via|localizzazione)")
    record["civic"] = label(r"(?:civico|n\.?\s*civico)")
    record["city"] = label(r"(?:comune|citt[aà])")
    record["province"] = label(r"(?:provincia|prov\.?|sigla\s+provincia)")
    record["brand"] = label(r"(?:marca|brand)")
    record["model"] = label(r"(?:modello|model)")
    record["serial"] = label(r"(?:matricola|seriale|serial(?:\s+number)?)")
    record["refrigerant"] = label(r"(?:gas|refrigerante|miscela)")
    record["charge_kg"] = label(r"(?:quantit[aà]|carico|precarica)(?:\s*\(?(?:in\s+)?kg\)?)?")
    record["purchase_reference"] = label(r"(?:fattura|scontrino|ordine|riferimento\s+acquisto|nr\.?\s*documento)")
    record["purchase_date"] = label(r"(?:data\s+(?:fattura|scontrino|acquisto|ordine|emissione))")
    record["intervention_date"] = label(r"(?:data\s+(?:di\s+)?(?:intervento|i?stallazione|installazione|dichiarazione\s+di\s+conformit[aà]))")
    record["installer_name"] = label(r"(?:installatore|nome\s+installatore)")
    record["installer_cf"] = label(r"(?:codice\s+fiscale\s+installatore|cf\s+installatore)")

    # Common Tecnomat/Bricoman order layout. Keep this narrow so store details
    # are not mistaken for customer details.
    if not record.get("client_name"):
        m = re.search(r"(?is)\bSignore\s+([A-ZÀ-ÖØ-Ý][A-ZÀ-ÖØ-Ý' -]{2,80}?)(?=\s+TECNOMAT\b|\s+VIA\b)", text)
        if m:
            record["client_name"] = " ".join(m.group(1).split()).title()
            record["client_type"] = "PRIVATO"
    address_match = re.search(
        r"(?im)\b((?:VIA|VIALE|CORSO|PIAZZA|PIAZZALE)\s+[A-ZÀ-ÖØ-Ý' .-]{2,80}?)\s+(\d+[A-Z]?)\s+(\d{5})\s+(MILANO)\b",
        text.upper(),
    )
    if address_match:
        record["install_address"] = " ".join(address_match.group(1).split()).title()
        record["civic"] = address_match.group(2)
        record["city"] = address_match.group(4).title()
        record["province"] = "MI"
    if not record.get("purchase_reference"):
        candidates = re.findall(r"(?i)\bORDINE\s+N[°º]?\s*(\d{6,14})\b", text)
        if candidates:
            best = max(candidates, key=lambda value: (len(value), candidates.index(value)))
            record["purchase_reference"] = f"ORDINE {best}"
    if not record.get("purchase_date"):
        m = re.search(r"(?i)\bData\s+emissione\s*:\s*(\d{1,2}/\d{1,2}/\d{4})\b", text)
        if m:
            record["purchase_date"] = m.group(1)
    if not record.get("intervention_date"):
        m = re.search(
            r"(?i)\bData\s+(?:di\s+)?(?:i?stallazione|installazione)\s*[:\-]?\s*(\d{1,2}/\d{1,2}/\d{4})\b",
            text,
        )
        if m:
            record["intervention_date"] = m.group(1)
    if not record.get("model"):
        m = re.search(r"\b(CL5000M\s+41/2\s+E)\b", text, re.I)
        if m or "4062321592103" in text or re.search(r"(?i)\bCOND\s+BOSCH\s+CL5000M\s+41\s+DUAL\b", text):
            record["model"] = "CL5000M 41/2 E"
    if record.get("model") and str(record["model"]).upper() == "CL5000M 41/2 E":
        record.setdefault("brand", "Bosch")
        if "TECNOMAT" in text.upper() and record.get("purchase_reference"):
            record.setdefault("sold_by_installer", False)

    cfs = CF_RE.findall(text.upper())
    if not record.get("client_cf") and cfs:
        record["client_cf"] = cfs[0]
    if not record.get("installer_cf") and len(cfs) > 1:
        record["installer_cf"] = cfs[-1]
    if not record.get("client_email"):
        m = EMAIL_RE.search(text)
        if m:
            record["client_email"] = m.group(0)
    # Never infer a customer phone from arbitrary numeric strings: orders,
    # product codes and store phone numbers are common in purchase documents.
    if not _usable_refrigerant(record.get("refrigerant")):
        m = REFRIGERANT_RE.search(text)
        if m:
            record["refrigerant"] = m.group(1)
    model_key = str(record.get("model") or "").upper().strip()
    if not _usable_kg(record.get("charge_kg")) and model_key not in MODEL_DEFAULTS:
        m = re.search(r"(?i)\bR-?32\b.{0,40}?\bKG\s*(\d{1,3}(?:[.,]\d{1,3})?)\b", text)
        if m:
            record["charge_kg"] = m.group(1)
    if not record.get("purchase_date"):
        dates = DATE_RE.findall(text)
        if dates:
            record["purchase_date"] = dates[0]

    # Serial candidates: prefer known manufacturer structures. Do not promote
    # arbitrary product/barcode strings into serials.
    hisense = re.findall(r"\b1K[A-Z0-9]{21}\b", text.upper())
    if hisense:
        record["serial"] = hisense[0]
    if model_key == "CL5000M 41/2 E":
        # Bosch labels use 86DM-XXX-NNNNNN-<product code>. OCR commonly reads
        # the D as 0/O and the final 2 as '>'. Repair only when the product code
        # is the model's verified Bosch code.
        bosch = re.search(r"\b[8B]6[DO0]M[- ](\d{3})[- ](\d{6})[- ](773370193[2>])", text.upper())
        if bosch:
            suffix = bosch.group(3).replace(">", "2")
            if suffix == MODEL_PRODUCT_CODES[model_key]:
                record["serial"] = f"86DM-{bosch.group(1)}-{bosch.group(2)}-{suffix}"
    bad_barcode = re.fullmatch(r"CS\d{10,}[A-Z0-9]*", str(record.get("serial") or ""), re.I)
    if bad_barcode:
        record["serial"] = ""

    if overrides:
        for key, value in overrides.items():
            if value not in (None, ""):
                record[key] = value
    return normalize_record({k: v for k, v in record.items() if v not in (None, "")})


def validate_record(record: Mapping[str, Any]) -> ValidationResult:
    normalized = normalize_record(record)
    warnings: list[str] = []
    missing = [key for key in REQUIRED_FOR_RENDER if _missing_value(normalized, key)]

    for key in ("client_cf", "installer_cf"):
        value = str(normalized.get(key) or "")
        if value and not CF_RE.fullmatch(value):
            warnings.append(f"{key}:format_unverified")
    serial = str(normalized.get("serial") or "")
    model = str(normalized.get("model") or "")
    brand = str(normalized.get("brand") or "")
    if brand.casefold() == "hisense" and model == "2AMW42U4RGC" and serial:
        if not (serial.startswith("1K") and len(serial) == 23):
            warnings.append("serial:hisense_expected_1K_23_chars")
    if normalized.get("refrigerant") and not REFRIGERANT_RE.fullmatch(str(normalized["refrigerant"])):
        warnings.append("refrigerant:format_unverified")
    if normalized.get("charge_kg") and not re.fullmatch(r"\d{1,3}(?:[.,]\d{1,3})?", str(normalized["charge_kg"])):
        warnings.append("charge_kg:format_unverified")
    for key in ("purchase_date", "intervention_date"):
        if normalized.get(key) and not _valid_date(str(normalized[key])):
            warnings.append(f"{key}:format_unverified")
    if normalized.get("gas_added") is True:
        for key in ("added_refrigerant", "added_kg", "added_gas_type"):
            if not normalized.get(key):
                missing.append(key)
    if normalized.get("gas_recovered") is True:
        for key in ("recovered_refrigerant", "recovered_kg"):
            if not normalized.get(key):
                missing.append(key)
    return ValidationResult(not missing, tuple(dict.fromkeys(missing)), tuple(warnings), normalized)


def render_docx(record: Mapping[str, Any], template: str | Path, output: str | Path) -> Path:
    from docx import Document

    validation = validate_record(record)
    r = validation.normalized
    template = Path(template)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(template))
    if len(doc.tables) < 6:
        raise ValueError("fgas_template_layout_unrecognized")

    t0, t1, t2, t3, t4, t5 = doc.tables[:6]
    _set(t0, 0, 2, str(r.get("client_name") or ""))
    _set(t0, 1, 2, f"CODICE FISCALE:{r.get('client_cf') or ''}")
    _set(t0, 2, 2, f"MAIL cliente:{r.get('client_email') or ''}")
    client_type = str(r.get("client_type") or "").upper()
    if client_type in {"PRIVATO", "IMPRESA", "ENTE", "ALTRO"}:
        row = {"PRIVATO": 1, "IMPRESA": 2, "ENTE": 3, "ALTRO": 4}[client_type]
        _set(t0, row, 1, f"{client_type} X")

    _set(t1, 1, 2, str(r.get("serial") or ""))
    _set(t1, 2, 2, str(r.get("brand") or ""))
    _set(t1, 3, 2, str(r.get("model") or ""))
    _set(t1, 4, 2, f"TIPO DI GAS:{r.get('refrigerant') or ''}")
    _set(t1, 4, 4, f"KG:{r.get('charge_kg') or ''}")
    if r.get("compressor_count"):
        _set(t1, 5, 4, str(r["compressor_count"]))
    _mark_yes_no(t1, 6, r.get("fixed_leak_detection"))
    _mark_yes_no(t1, 7, r.get("hermetically_sealed"))
    equipment = str(r.get("equipment_type") or "").upper()
    if equipment:
        mapping = {
            "POMPA DI CALORE FISSA": (9, 1),
            "FISSA DI CONDIZIONAMENTO": (9, 3),
            "FISSA DI REFRIGERAZIONE": (9, 5),
        }
        if equipment in mapping:
            row, col = mapping[equipment]
            _set(t1, row, col, equipment + " X")

    _set(t2, 0, 2, str(r.get("install_address") or ""))
    _set(t2, 1, 1, "CIVICO:")
    _set(t2, 1, 2, str(r.get("civic") or ""))
    _set(t2, 2, 1, "COMUNE:")
    _set(t2, 2, 2, str(r.get("city") or ""))
    _set(t2, 2, 3, f"PROVINCIA:{r.get('province') or ''}")
    _set(t2, 3, 1, f"DESTINAZIONE D'USO: (E1,E2,...) {r.get('use_destination') or ''}")

    _set(t3, 0, 1, f"DATA INTERVENTO/DICHIARAZIONE DI CONFORMITÀ (OBBLIGATORIO!): {r.get('intervention_date') or ''}")
    sold = r.get("sold_by_installer")
    if sold is True:
        _set(t3, 1, 2, "SI X (indicare data e nr vostra fattura lavoro)")
    elif sold is False:
        _set(t3, 1, 3, "NO X (indicare data e nr scontrino acquisto impianto da parte del cliente)")
    _set(t3, 2, 2, str(r.get("purchase_reference") or ""))
    _set(t3, 3, 2, str(r.get("purchase_date") or ""))
    _set(t3, 4, 1, f"NOME E COGNONE DELL'INSTALLATORE:{r.get('installer_name') or ''}")
    _set(t3, 5, 1, f"CODICE FISCALE INSTALLATORE:{r.get('installer_cf') or ''}")

    _set(t4, 0, 3, f"MISCELA (Es. R32): {r.get('refrigerant') or ''}")
    _set(t4, 0, 5, f"QUANTITÀ (IN KG):{r.get('charge_kg') or ''}")
    recovered = r.get("gas_recovered")
    if recovered is not None:
        _set(t4, 2, 1, f"HAI RECUPERATO GAS? {'SI' if recovered else 'NO'}")
    if recovered:
        _set(t4, 2, 3, f"MISCELA (Es. R32): {r.get('recovered_refrigerant') or ''}")
        _set(t4, 2, 5, f"QUANTITÀ (IN KG):{r.get('recovered_kg') or ''}")
    added = r.get("gas_added")
    if added is not None:
        _set(t4, 3, 2, f"HAI AGGIUNTO GAS?: {'SI' if added else 'NO'}  MISCELA (Es. R32): {r.get('added_refrigerant') or ''}  QUANTITÀ (IN KG): {r.get('added_kg') or ''}")
    gas_type = str(r.get("added_gas_type") or "").upper()
    if added and gas_type in {"VERGINE", "RIGENERATO", "RICICLATO"}:
        col = {"VERGINE": 3, "RIGENERATO": 4, "RICICLATO": 6}[gas_type]
        _set(t4, 4, col, gas_type + " X")
    supplied = r.get("added_gas_client_supplied")
    if added and supplied is not None:
        _set(t4, 5, 3 if supplied else 5, ("SI" if supplied else "NO") + " X")

    _set(t5, 0, 1, str(r.get("observations") or "")[:300])
    doc.save(str(output))
    return output


def convert_docx_to_pdf(docx_path: str | Path, pdf_path: str | Path) -> Path:
    docx_path = Path(docx_path)
    pdf_path = Path(pdf_path)
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not soffice:
        return _render_simple_pdf(docx_path, pdf_path)
    completed = subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)],
        stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=90, check=False,
    )
    generated = pdf_path.parent / (docx_path.stem + ".pdf")
    if completed.returncode == 0 and generated.exists():
        if generated != pdf_path:
            generated.replace(pdf_path)
        return pdf_path
    return _render_simple_pdf(docx_path, pdf_path)


def _render_simple_pdf(docx_path: Path, pdf_path: Path) -> Path:
    from docx import Document
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    doc = Document(str(docx_path))
    lines: list[str] = ["MODULO INSTALLAZIONE - F-Gas"]
    for table in doc.tables:
        for row in table.rows:
            values = [" ".join(cell.text.split()) for cell in row.cells]
            compact = " | ".join(value for index, value in enumerate(values) if value and value not in values[:index])
            if compact:
                lines.append(compact)
    c = canvas.Canvas(str(pdf_path), pagesize=A4)
    width, height = A4
    y = height - 40
    c.setFont("Helvetica", 8)
    for line in lines:
        chunks = [line[i:i+125] for i in range(0, len(line), 125)] or [""]
        for chunk in chunks:
            if y < 40:
                c.showPage(); c.setFont("Helvetica", 8); y = height - 40
            c.drawString(30, y, chunk); y -= 11
    c.save()
    return pdf_path


def _set(table: Any, row: int, col: int, value: str) -> None:
    cell = table.cell(row, col)
    cell.text = value


def _mark_yes_no(table: Any, row: int, value: Any) -> None:
    if value is True:
        _set(table, row, 4, "SI X")
        _set(table, row, 6, "NO")
    elif value is False:
        _set(table, row, 4, "SI")
        _set(table, row, 6, "NO X")


def _missing_value(record: Mapping[str, Any], key: str) -> bool:
    if key not in record or record[key] is None:
        return True
    value = record[key]
    if isinstance(value, bool):
        return False
    if isinstance(value, str):
        return not value.strip()
    return False


def _valid_date(value: str) -> bool:
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y"):
        try:
            datetime.strptime(value, fmt)
            return True
        except ValueError:
            pass
    return False


def _usable_refrigerant(value: Any) -> bool:
    return bool(value and REFRIGERANT_RE.search(str(value)))


def _usable_kg(value: Any) -> bool:
    return bool(value and re.search(r"\d", str(value)))


__all__ = [
    "MODEL_DEFAULTS", "REQUIRED_FOR_RENDER", "ValidationResult",
    "convert_docx_to_pdf", "extract_record", "normalize_record", "render_docx", "validate_record",
]
