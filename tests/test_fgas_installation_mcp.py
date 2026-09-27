from __future__ import annotations

from pathlib import Path

from docx import Document

from ralfloop_agent.fgas_installation import extract_record, render_docx, validate_record
from scripts.ralf_fgas_installation_mcp_server import FGasMCPServer, TOOLS, _parse_drive_files


COMPLETE = {
    "client_name": "Mario Rossi",
    "client_cf": "RSSMRA80A01F205X",
    "client_email": "mario@example.org",
    "client_type": "PRIVATO",
    "installer_name": "Installatore Demo",
    "installer_cf": "BNCLGU80A01F205Z",
    "install_address": "Via Esempio",
    "civic": "10",
    "city": "Milano",
    "province": "MI",
    "use_destination": "E1 Casa",
    "brand": "Hisense",
    "model": "2AMW42U4RGC",
    "serial": "1KA0140012AP08BX5JB1924",
    "refrigerant": "R32",
    "charge_kg": "0,95",
    "compressor_count": "1",
    "fixed_leak_detection": False,
    "hermetically_sealed": True,
    "equipment_type": "POMPA DI CALORE FISSA",
    "intervention_date": "27/09/2026",
    "sold_by_installer": False,
    "purchase_reference": "SCONTRINO 123",
    "purchase_date": "20/09/2026",
    "gas_recovered": False,
    "gas_added": False,
}


def _template(path: Path) -> Path:
    doc = Document()
    for rows, cols in ((5, 3), (10, 7), (4, 4), (6, 4), (6, 7), (1, 2)):
        table = doc.add_table(rows=rows, cols=cols)
        for r in range(rows):
            for c in range(cols):
                table.cell(r, c).text = f"T{len(doc.tables)-1}R{r}C{c}"
    doc.save(path)
    return path


def test_extract_applies_known_model_defaults_and_serial_rule():
    text = """
    Cliente: Mario Rossi
    RSSMRA80A01F205X
    Marca: Hisense
    Modello: 2AMW42U4RGC
    Matricola: 1KA0140012AP08BX5JB1924
    Indirizzo: Via Esempio
    Comune: Milano
    Provincia: MI
    Data intervento: 27/09/2026
    Installatore: Installatore Demo
    Codice fiscale installatore: BNCLGU80A01F205Z
    """
    record = extract_record(text)
    assert record["refrigerant"] == "R32"
    assert record["charge_kg"] == "0,95"
    assert record["serial"] == "1KA0140012AP08BX5JB1924"


def test_extract_rejects_known_false_barcode_as_serial():
    record = extract_record("Matricola: CS661476958030M2\nMarca: Hisense")
    assert not record.get("serial")


def test_validation_never_invents_missing_required_fields():
    result = validate_record({"client_name": "Mario Rossi"})
    assert result.ok is False
    assert "serial" in result.missing
    assert "installer_cf" in result.missing
    assert "serial" not in result.normalized


def test_render_populates_synthetic_template(tmp_path: Path):
    template = _template(tmp_path / "template.docx")
    output = tmp_path / "out.docx"
    render_docx(COMPLETE, template, output)
    doc = Document(output)
    assert doc.tables[0].cell(0, 2).text == "Mario Rossi"
    assert "RSSMRA80A01F205X" in doc.tables[0].cell(1, 2).text
    assert doc.tables[1].cell(1, 2).text == "1KA0140012AP08BX5JB1924"
    assert "R32" in doc.tables[1].cell(4, 2).text
    assert "0,95" in doc.tables[1].cell(4, 4).text
    assert "Milano" in doc.tables[2].cell(2, 2).text
    assert "27/09/2026" in doc.tables[3].cell(0, 1).text


def test_drive_parser_selects_real_docx_shape():
    text = """## Files (3)
1Oed_oGhMNeR4p2RS5-87uXacB5zSWoZ5 | MODULO INSTALLAZIONE 2.docx | vnd.openxmlformats-officedocument.wordprocessingml.document | Oct 2 | 238.8 KB
1DxhZvSlCeeeeU0YDdLXN755sUNSiDcIc | MODULO INSTALLAZIONE 2.docx.pdf | pdf | Jun 4 | 242.0 KB
"""
    rows = _parse_drive_files(text)
    from scripts.ralf_fgas_installation_mcp_server import DriveFGasSource
    assert DriveFGasSource.select_template(rows)["id"] == "1Oed_oGhMNeR4p2RS5-87uXacB5zSWoZ5"


def test_mcp_surface_is_semantic_and_strict():
    assert set(TOOLS) == {
        "fgas_drive_status", "fgas_extract_from_text", "fgas_validate_installation",
        "fgas_render_installation", "fgas_prepare_from_whatsapp",
    }
    serialized = str(TOOLS).casefold()
    for forbidden in ("selector", "xpath", "javascript", "shell", "delete", "share", "upload"):
        assert forbidden not in serialized
    assert all(schema["additionalProperties"] is False for schema in TOOLS.values())


class _FakeDrive:
    def status(self):
        return {"ok": True, "status": "READY", "template": {"id": "demo"}, "examples": [], "side_effects": 0}


class _FakeWhatsApp:
    def collect(self, **_kwargs):
        return {"ok": True, "status": "READY", "chat_id": "wa_chat_" + "a" * 16, "message_ids": [], "text": "Cliente: Mario Rossi", "media": [], "side_effects": 0}


def test_prepare_from_whatsapp_stops_on_missing_fields():
    server = FGasMCPServer(drive=_FakeDrive(), whatsapp=_FakeWhatsApp())
    result = server.call("fgas_prepare_from_whatsapp", {"chat_id": "wa_chat_" + "a" * 16})
    payload = result["structuredContent"]
    assert payload["status"] == "INCOMPLETE"
    assert "render" not in payload
    assert "serial" in payload["missing"]


def test_mcp_rejects_schema_violations_before_execution():
    server = FGasMCPServer(drive=_FakeDrive(), whatsapp=_FakeWhatsApp())
    bad_path = server.call("fgas_render_installation", {
        "record": COMPLETE, "output_basename": "../../escape",
    })
    bad_bool = server.call("fgas_render_installation", {
        "record": {**COMPLETE, "gas_added": "no"},
    })
    assert bad_path["structuredContent"]["status"] == "POLICY_DENIED"
    assert bad_bool["structuredContent"]["status"] == "POLICY_DENIED"


def test_live_catalog_config_registers_fgas_aliases():
    import json
    config = json.loads(Path("config/mcp_catalog_v1.json").read_text())
    providers = {row["id"]: row["socket"] for row in config["providers"]}
    assert providers["fgas"] == "/run/ralf-fgas-mcp/mcp.sock"
    assert "fgas" in config["term_aliases"]["condizionatore"]
    assert "fgas" in config["term_aliases"]["climatizzatore"]
